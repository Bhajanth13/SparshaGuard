"""Simple workbook: every value found in the papers, with its full reference on the same row.

Run:  python tools/build_paper_values_xlsx.py
"""
import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("src", HERE / "build_sources_xlsx.py")
src = importlib.util.module_from_spec(spec)
spec.loader.exec_module(src)
OUT = HERE.parent / "docs" / "SparshaGuard_paper_values.xlsx"

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A4D")
GROUP = PatternFill("solid", fgColor="DCE6EE")
side = Side(style="thin", color="C9D3DC")
BD = Border(left=side, right=side, top=side, bottom=side)
USED = {"In model v0.3": "Yes (in model)", "Pending integration": "Found, not added yet",
        "Benchmark only": "Accuracy benchmark", "Context only": "Background only",
        "Whole-body study": "Whole-body study (doc)",
        "Fusion study": "Fusion study (doc)"}


def ref_text(s):
    sid, auth, title, venue, year, typ, pref, url, note = s
    y = year if year else "n.d."
    return f'{auth} ({y}). "{title}". {venue}.'


src_by_id = {s[0]: s for s in src.SOURCES}
order = [s[0] for s in src.SOURCES]

wb = Workbook()
ws = wb.active
ws.title = "Values from Papers"
cols = ["#", "Parameter", "Value", "Unit", "Measured under / conditions", "Paper",
        "Full reference", "Where in the paper", "Link", "Used in our model?"]
widths = [5, 36, 24, 11, 34, 26, 60, 16, 42, 20]
for i, (c, w) in enumerate(zip(cols, widths), start=1):
    cell = ws.cell(row=1, column=i, value=c)
    cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    cell.fill, cell.border = HDR, BD
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[cell.column_letter].width = w
ws.row_dimensions[1].height = 28

r, n = 1, 0
for sid in order:
    rows = [v for v in src.V if v[5] == sid]
    if not rows:
        continue
    s = src_by_id[sid]
    r += 1
    g = ws.cell(row=r, column=1, value=f"{src.SHORT[sid]}  —  {s[3]}")
    g.font = Font(name=F, bold=True, size=10)
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).fill = GROUP
        ws.cell(row=r, column=c).border = BD
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
    for cat, par, val, unit, cond, _sid, loc, status, used, note in rows:
        r += 1
        n += 1
        cond_full = "; ".join(x for x in (cond, note) if x)
        vals = [n, par, val, unit, cond_full, src.SHORT[sid], ref_text(s), loc, s[7], USED[status]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = Font(name=F, size=10), BD
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        link = ws.cell(row=r, column=9)
        link.hyperlink = s[7]
        link.font = Font(name=F, size=10, color="0563C1", underline="single")
ws.auto_filter.ref = f"A1:J{r}"

rs = wb.create_sheet("Reference List")
for i, (c, w) in enumerate(zip(["#", "Reference", "Link", "Values taken"], [5, 100, 50, 12]), start=1):
    cell = rs.cell(row=1, column=i, value=c)
    cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    cell.fill, cell.border = HDR, BD
    rs.column_dimensions[cell.column_letter].width = w
for i, sid in enumerate(order, start=1):
    s = src_by_id[sid]
    cnt = sum(1 for v in src.V if v[5] == sid)
    for c, v in enumerate([i, ref_text(s), s[7], cnt], start=1):
        cell = rs.cell(row=i + 1, column=c, value=v)
        cell.font, cell.border = Font(name=F, size=10), BD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    lc = rs.cell(row=i + 1, column=3)
    lc.hyperlink = s[7]
    lc.font = Font(name=F, size=10, color="0563C1", underline="single")

wb.save(OUT)
print("wrote", OUT, n, "values from", len(order), "papers")
